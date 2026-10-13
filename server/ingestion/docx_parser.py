from pathlib import Path
from docx import Document


def parse_docx(file_path: str) -> dict:
    """
    Extract text from a DOCX document.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    document = Document(file_path)

    text_parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract table content
    for table_index, table in enumerate(document.tables):

        text_parts.append(
            f"\n[Table {table_index + 1}]"
        )

        for row in table.rows:

            row_values = [
                cell.text.strip()
                for cell in row.cells
            ]

            text_parts.append(
                " | ".join(row_values)
            )

    return {
        "text": "\n".join(text_parts).strip(),
        "metadata": {
            "filename": path.name,
            "file_type": "docx",
            "source": str(path)
        }
    }